"""
tests/report.py

Formateo del informe en consola y Excel.
No tocar al añadir nuevos test cases.
"""
from __future__ import annotations
import json
import os
from dataclasses import dataclass, field
from datetime import datetime


@dataclass
class TurnReport:
    turn_number:  int
    user_message: str
    reply:        str
    stage:        str
    passed:       bool
    failures:     list[str] = field(default_factory=list)


@dataclass
class CaseReport:
    case_id:     str
    case_type:   str
    description: str
    file_path:   str
    passed:      bool
    turns:       list[TurnReport] = field(default_factory=list)
    error:       str | None = None
    duration_ms: int = 0


# ------------------------------------------------------------------ #
#  Consola                                                             #
# ------------------------------------------------------------------ #

def print_results(reports: list[CaseReport]) -> None:
    total  = len(reports)
    passed = sum(1 for r in reports if r.passed)
    failed = total - passed
    ms     = sum(r.duration_ms for r in reports)

    _sep()
    print("  DENTAL AGENT — TEST SUITE")
    print(f"  {total} caso(s)  ·  {datetime.now().strftime('%H:%M:%S')}\n")

    for r in reports:
        icon = "✅" if r.passed else "❌"
        print(f"  {icon}  {r.case_id:<12} {r.description}")

        if not r.passed:
            if r.error:
                print(f"               ⚠  Error: {r.error}")
            for t in r.turns:
                if not t.passed:
                    print(f"               Turno {t.turn_number}  user: \"{t.user_message}\"")
                    for f in t.failures:
                        print(f"                 · {f}")
                    preview = t.reply[:120] + ("..." if len(t.reply) > 120 else "")
                    print(f"               Reply: \"{preview}\"")

    _sep()
    print(f"  ✅ {passed} pasados   ❌ {failed} fallidos   ⏱  {ms / 1000:.1f}s")
    _sep()
    print()


def _sep():
    print("═" * 58)


# ------------------------------------------------------------------ #
#  Excel                                                               #
# ------------------------------------------------------------------ #

def update_inventory(reports: list[CaseReport], output_path: str = "tests/inventory.xlsx") -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  ⚠  openpyxl no instalado → pip install openpyxl")
        return

    wb = Workbook()

    # ================================================================ #
    #  Hoja 1: inventario por caso                                      #
    # ================================================================ #
    ws = wb.active
    ws.title = "Test Inventory"

    HEADERS = [
        "#", "TYPE", "ID TEST", "DESCRIPTION", "FILE PATH", "TAGS",
        "USER MESSAGES", "AGENT REPLIES", "LAST RESULT", "LAST RUN", "OBSERVATIONS",
    ]
    WIDTHS  = [
        4,   12,     12,       44,            34,          18,
        48,              64,              13,            20,         32,
    ]

    styles = _excel_styles(Font, PatternFill, Alignment, Border, Side)
    _setup_sheet(ws, HEADERS, WIDTHS, styles, get_column_letter)

    now = datetime.now().strftime("%Y-%m-%d %H:%M")

    for row_idx, r in enumerate(reports, 2):
        tags   = ", ".join(_read_tags(r.file_path))
        result = "PASS" if r.passed else "FAIL"
        obs    = r.error or _first_failure(r)

        # Nuevo: guardamos en el inventario el texto enviado y la respuesta recibida.
        user_messages = _format_case_messages(r)
        agent_replies = _format_case_replies(r)

        row_fill = styles["pass_fill"] if r.passed else styles["fail_fill"]

        values = [
            row_idx - 1, r.case_type, r.case_id, r.description, r.file_path,
            tags, user_messages, agent_replies, result, now, obs,
        ]

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill   = row_fill
            cell.border = styles["border"]
            cell.alignment = styles["center"] if col in (1, 2, 9) else styles["left"]

            if col == 9:
                cell.font = styles["pass_font"] if r.passed else styles["fail_font"]
            else:
                cell.font = styles["normal_font"]

        # Más altura porque ahora hay conversaciones completas en la fila.
        ws.row_dimensions[row_idx].height = max(36, min(160, 20 + 18 * max(1, len(r.turns))))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    # ================================================================ #
    #  Hoja 2: detalle por turno                                        #
    # ================================================================ #
    ws_turns = wb.create_sheet("Turn Details")
    TURN_HEADERS = [
        "ID TEST", "TURN", "USER MESSAGE", "AGENT REPLY", "STAGE",
        "RESULT", "FAILURES", "FILE PATH",
    ]
    TURN_WIDTHS = [12, 8, 56, 72, 24, 12, 48, 34]
    _setup_sheet(ws_turns, TURN_HEADERS, TURN_WIDTHS, styles, get_column_letter)

    row_idx = 2
    for r in reports:
        if not r.turns:
            values = [
                r.case_id, "", "", "", "", "FAIL" if not r.passed else "PASS",
                r.error or "Sin turnos ejecutados", r.file_path,
            ]
            _write_turn_row(ws_turns, row_idx, values, r.passed, styles)
            row_idx += 1
            continue

        for t in r.turns:
            values = [
                r.case_id,
                t.turn_number,
                t.user_message,
                t.reply,
                t.stage,
                "PASS" if t.passed else "FAIL",
                "\n".join(t.failures),
                r.file_path,
            ]
            _write_turn_row(ws_turns, row_idx, values, t.passed, styles)
            ws_turns.row_dimensions[row_idx].height = max(36, min(180, 18 * max(2, _line_count(t.user_message), _line_count(t.reply))))
            row_idx += 1

    ws_turns.freeze_panes = "A2"
    ws_turns.auto_filter.ref = f"A1:{get_column_letter(len(TURN_HEADERS))}1"

    wb.save(output_path)
    print(f"  📊 Inventario → {output_path}\n")


def _excel_styles(Font, PatternFill, Alignment, Border, Side) -> dict:
    thin        = Side(style="thin", color="CCCCCC")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)
    h_fill      = PatternFill("solid", start_color="1F4E79")
    h_font      = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    center      = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left        = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    pass_fill   = PatternFill("solid", start_color="E2EFDA")
    fail_fill   = PatternFill("solid", start_color="FCE4D6")
    pass_font   = Font(bold=True, color="375623", name="Calibri", size=10)
    fail_font   = Font(bold=True, color="9C0006", name="Calibri", size=10)
    normal_font = Font(name="Calibri", size=10)

    return {
        "thin": thin,
        "border": border,
        "h_fill": h_fill,
        "h_font": h_font,
        "center": center,
        "left": left,
        "pass_fill": pass_fill,
        "fail_fill": fail_fill,
        "pass_font": pass_font,
        "fail_font": fail_font,
        "normal_font": normal_font,
    }


def _setup_sheet(ws, headers: list[str], widths: list[int], styles: dict, get_column_letter) -> None:
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = styles["h_font"]
        cell.fill = styles["h_fill"]
        cell.alignment = styles["center"]
        cell.border = styles["border"]
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 28


def _write_turn_row(ws, row_idx: int, values: list, passed: bool, styles: dict) -> None:
    row_fill = styles["pass_fill"] if passed else styles["fail_fill"]

    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col, value=val)
        cell.fill = row_fill
        cell.border = styles["border"]
        cell.alignment = styles["center"] if col in (2, 6) else styles["left"]

        if col == 6:
            cell.font = styles["pass_font"] if passed else styles["fail_font"]
        else:
            cell.font = styles["normal_font"]


def _format_case_messages(r: CaseReport) -> str:
    return "\n".join(f"T{t.turn_number}: {t.user_message}" for t in r.turns)


def _format_case_replies(r: CaseReport) -> str:
    return "\n".join(f"T{t.turn_number}: {t.reply}" for t in r.turns)


def _line_count(text: str) -> int:
    if not text:
        return 1
    # Aproximación simple para que las respuestas largas sean legibles en Excel.
    explicit_lines = text.count("\n") + 1
    wrapped_lines = max(1, len(text) // 80)
    return explicit_lines + wrapped_lines


def _read_tags(file_path: str) -> list[str]:
    try:
        with open(file_path) as f:
            return json.load(f).get("tags", [])
    except Exception:
        return []


def _first_failure(r: CaseReport) -> str:
    for t in r.turns:
        if t.failures:
            return f"T{t.turn_number}: {t.failures[0]}"
    return ""
